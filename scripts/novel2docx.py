# -*- coding: utf-8 -*-
"""把小说 Markdown 转成中文排版规范的 .docx。

用法:
    python novel2docx.py 输入.md [输出.docx]

排版规则(中文小说/网文):
    - 书名:黑体 三号 加粗 居中
    - 章标题:黑体 小三 加粗 居中
    - 正文:宋体 小四(12pt),首行缩进 2 字符,两端对齐,1.5 倍行距,段间不空行
    - 西文/数字用 Times New Roman
"""
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

EAST_FONT = "宋体"
EAST_FONT_HEAD = "黑体"
ASCII_FONT = "Times New Roman"


def _set_font(run, ascii_font, east_font, size_pt, bold=False):
    run.font.name = ascii_font
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), ascii_font)
    rFonts.set(qn("w:hAnsi"), ascii_font)
    rFonts.set(qn("w:eastAsia"), east_font)


def _first_line_chars(paragraph, chars=200):
    """首行缩进 N 个全角字符(200 = 2 字),随字号缩放。"""
    pPr = paragraph._p.get_or_add_pPr()
    ind = pPr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        pPr.append(ind)
    ind.set(qn("w:firstLineChars"), str(chars))
    ind.attrib.pop(qn("w:firstLine"), None)


def _strip_inline(text):
    return re.sub(r"(\*\*|__|`|~~)", "", text).strip()


def main():
    if len(sys.argv) < 2:
        sys.exit("用法: python novel2docx.py 输入.md [输出.docx]")
    in_path = sys.argv[1]
    out_path = sys.argv[2] if len(sys.argv) > 2 else re.sub(r"\.md$", ".docx", in_path)

    with open(in_path, encoding="utf-8") as f:
        lines = f.read().splitlines()

    doc = Document()
    # 正文默认样式兜底
    normal = doc.styles["Normal"]
    normal.font.name = ASCII_FONT
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), EAST_FONT)

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        if stripped in ("---", "***"):
            continue

        if stripped.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(12)
            _set_font(p.add_run(_strip_inline(stripped[2:])), ASCII_FONT, EAST_FONT_HEAD, 16, bold=True)
        elif stripped.startswith("## "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            _set_font(p.add_run(_strip_inline(stripped[3:])), ASCII_FONT, EAST_FONT_HEAD, 15, bold=True)
        elif stripped.startswith("### "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(3)
            _set_font(p.add_run(_strip_inline(stripped[4:])), ASCII_FONT, EAST_FONT_HEAD, 14, bold=True)
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf = p.paragraph_format
            pf.line_spacing = 1.5
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
            _first_line_chars(p, 200)
            _set_font(p.add_run(_strip_inline(stripped)), ASCII_FONT, EAST_FONT, 12)

    doc.save(out_path)
    print("已生成:", out_path)


if __name__ == "__main__":
    main()
